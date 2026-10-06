import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from app.images import prepare_for_slide, title_background
from app.models import AnnotationOverrides, AnnotationSettings, TitlePhoto
from app.presentation import (
    SLIDE_H, SLIDE_W, annotation_lines, build_slide_plan, fit_rect, format_date, write_pptx,
)
from app.project import add_observations
from tests.conftest import make_obs
from tests.helpers import norm, project, source

URL = "https://www.inaturalist.org/observations?taxon_id=47170&place_id=1"


def _jpeg(path: Path, w: int, h: int, color=(120, 60, 30), exif_orientation=None) -> Path:
    im = Image.new("RGB", (w, h), color)
    kwargs = {"quality": 85}
    if exif_orientation:
        exif = Image.Exif()
        exif[0x0112] = exif_orientation
        kwargs["exif"] = exif
    im.save(path, "JPEG", **kwargs)
    return path


def _deck(settings_kw=None, sources=None):
    obs = norm(
        make_obs(1, 47692, faves=17, user="alice", name="Alice Smith", observed_on="2026-11-14",
                 place="Point Reyes, CA, US", photos=((101, 4000, 3000), (102, 2000, 3000), (103, 3000, 3000))),
        make_obs(2, 48460, faves=0, user="bob", photos=((201, 1600, 900),)),  # identified only as Life
        make_obs(3, 47602, faves=2, user="carol", photos=((301, 1000, 1000),)),
    )
    pr = project(*(sources or [source("a", "alice")]), **(settings_kw or {}))
    pr = add_observations(pr, [1, 2, 3], obs, {1: ["a"], 2: ["a"], 3: ["a"]}, placement="append")
    return pr, obs


def _images(tmp_path, obs):
    out = {}
    for o in obs.values():
        for p in o["photos"]:
            src = _jpeg(tmp_path / f"{p['id']}.jpg", p["width"], p["height"])
            out[p["id"]] = prepare_for_slide(src, tmp_path / "work")
    return out


# ---------------------------------------------------------------- plan

def test_plan_multiple_photos_selection_and_life():
    pr, obs = _deck()
    st = {s.id: s for s in pr.observations}
    st[1] = st[1].model_copy(update={"selected_photo_ids": [101, 103]})  # 3 photos, 2 selected
    pr = pr.model_copy(update={"observations": list(st.values())})
    plan = build_slide_plan(pr, obs)
    kinds = [s["kind"] for s in plan["slides"]]
    assert kinds == ["title", "image", "image", "image", "image"]
    imgs = plan["slides"][1:]
    assert [s["photo_id"] for s in imgs] == [101, 103, 201, 301]
    assert imgs[0]["observation_id"] == imgs[1]["observation_id"] == 1
    assert imgs[2]["lines"] == []  # "Life" -> no taxon annotation
    assert imgs[0]["lines"][0] == {"field": "scientific", "text": "Hygrocybe conica", "italic": True, "size": 28}
    assert plan["counts"] == {"total": 5, "title": 1, "dividers": 0, "images": 4, "observations": 3}


def test_plan_min_faves_dividers_and_order():
    pr, obs = _deck({"grouping": "family", "dividers": True, "min_faves": 1})
    pr = pr.model_copy(update={"order": [3, 1, 2]})
    plan = build_slide_plan(pr, obs)
    seq = [(s["kind"], s.get("text") or s.get("photo_id")) for s in plan["slides"][1:]]
    assert seq == [("divider", "Amanitaceae"), ("image", 301), ("divider", "Hygrophoraceae"),
                   ("image", 101), ("image", 102), ("image", 103)]
    assert plan["slides"][1]["rank_label"] == "Family"


def test_plan_skips_unavailable_and_flags_deleted_photos():
    pr, obs = _deck()
    st = {s.id: s for s in pr.observations}
    st[3] = st[3].model_copy(update={"status": "unavailable"})
    st[1] = st[1].model_copy(update={"selected_photo_ids": [101, 999]})
    pr = pr.model_copy(update={"observations": list(st.values())})
    plan = build_slide_plan(pr, obs)
    assert [s.get("photo_id") for s in plan["slides"][1:]] == [101, 201]
    assert any("no longer available" in w for w in plan["warnings"])
    assert any("deleted" in w for w in plan["warnings"])


def test_annotation_fields_overrides_and_observer_rules():
    on = AnnotationSettings(scientific=True, common_name=True, date=True, location=True)
    pr, obs = _deck({"annotations": on})
    st = {s.id: s for s in pr.observations}
    lines = annotation_lines(obs[1], st[1], pr)
    assert [l["text"] for l in lines] == ["Hygrocybe conica", "Witch's Hat", "November 14, 2026", "Point Reyes, CA, US"]
    assert [l["italic"] for l in lines] == [True, False, False, False]
    # username-only project: no observer line unless asked for
    pr_obs = pr.model_copy(update={"settings": pr.settings.model_copy(update={"annotations": on.model_copy(update={"observer": True})})})
    assert annotation_lines(obs[1], st[1], pr_obs)[-1]["text"] == "Photo: Alice Smith"
    assert annotation_lines(obs[3], st[3], pr_obs)[-1]["text"] == "Photo: carol"  # login fallback
    # URL project: observer on by default
    url_pr, _ = _deck(sources=[source("a", URL)])
    assert annotation_lines(obs[1], st[1], url_pr)[-1]["text"] == "Photo: Alice Smith"
    # user edits win; "" hides a line
    edited = st[1].model_copy(update={"overrides": AnnotationOverrides(scientific="Hygrocybe sp. nov.", location="", date="Fall 2026")})
    texts = [l["text"] for l in annotation_lines(obs[1], edited, pr)]
    assert texts == ["Hygrocybe sp. nov.", "Witch's Hat", "Fall 2026"]


def test_format_date():
    assert format_date("2026-11-14") == "November 14, 2026"
    assert format_date("2026-01-01") == "January 1, 2026"
    assert format_date(None) == "" and format_date("garbage") == ""


@pytest.mark.parametrize("w,h", [(4000, 3000), (2000, 3000), (3000, 3000), (1920, 1080), (6000, 1000), (500, 4000)])
def test_fit_rect_never_crops_and_centers(w, h):
    left, top, fw, fh = fit_rect(w, h)
    assert 0 <= left and 0 <= top
    assert left + fw <= SLIDE_W and top + fh <= SLIDE_H
    assert fw == SLIDE_W or fh == SLIDE_H  # fills one dimension
    assert abs(fw / fh - w / h) < 0.01  # aspect preserved
    assert abs(left - (SLIDE_W - fw - left)) <= 1 and abs(top - (SLIDE_H - fh - top)) <= 1


# ---------------------------------------------------------------- pptx

def _build(tmp_path, pr, obs, title=None):
    plan = build_slide_plan(pr, obs)
    out = tmp_path / "deck.pptx"
    write_pptx(plan, _images(tmp_path, obs), title, out)
    return Presentation(str(out)), out, plan


def _pictures(slide):
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_pptx_is_16x9_with_uncropped_centered_photos(tmp_path):
    pr, obs = _deck({"title": "Fungi of Marin", "presenter": "Alan"})
    prs, out, plan = _build(tmp_path, pr, obs)
    assert prs.slide_width == Emu(12192000) and prs.slide_height == Emu(6858000)
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 1e-3
    assert len(prs.slides) == len(plan["slides"]) == 6
    dims = {101: (4000, 3000), 102: (2000, 3000), 103: (3000, 3000), 201: (1600, 900), 301: (1000, 1000)}
    for slide, spec in zip(list(prs.slides)[1:], plan["slides"][1:]):
        (pic,) = _pictures(slide)
        w, h = dims[spec["photo_id"]]
        assert pic.crop_left == pic.crop_right == pic.crop_top == pic.crop_bottom == 0
        assert abs(pic.width / pic.height - w / h) < 0.01
        assert pic.left >= 0 and pic.top >= 0
        assert pic.left + pic.width <= prs.slide_width and pic.top + pic.height <= prs.slide_height
        assert abs(pic.left - (prs.slide_width - pic.left - pic.width)) <= 2
        assert abs(pic.top - (prs.slide_height - pic.top - pic.height)) <= 2
        assert pic.image.size == (w, h)  # full resolution embedded


def test_photos_embedded_byte_for_byte_and_stored(tmp_path):
    pr, obs = _deck()
    imgs = _images(tmp_path, obs)
    plan = build_slide_plan(pr, obs)
    out = tmp_path / "deck.pptx"
    write_pptx(plan, imgs, None, out)
    with zipfile.ZipFile(out) as z:
        media = [i for i in z.infolist() if i.filename.startswith("ppt/media/")]
        assert len(media) == 5
        assert all(i.compress_type == zipfile.ZIP_STORED for i in media)
        originals = {imgs[pid].path.read_bytes() for pid in imgs}
        assert {z.read(i) for i in media} == originals  # no re-encoding


@pytest.mark.parametrize("bg,rgb", [("black", "000000"), ("white", "FFFFFF")])
def test_backgrounds(tmp_path, bg, rgb):
    pr, obs = _deck({"background": bg})
    prs, _, _ = _build(tmp_path, pr, obs)
    for slide in prs.slides:
        assert str(slide.background.fill.fore_color.rgb) == rgb


def test_annotation_formatting(tmp_path):
    pr, obs = _deck({"annotations": AnnotationSettings(common_name=True)})
    prs, _, _ = _build(tmp_path, pr, obs)
    slide = prs.slides[1]
    box = next(s for s in slide.shapes if s.has_text_frame and s.name == "Annotation")
    paras = box.text_frame.paragraphs
    sci, common = paras[0].runs[0], paras[1].runs[0]
    assert sci.text == "Hygrocybe conica" and sci.font.italic is True and sci.font.size.pt == 28
    assert str(sci.font.color.rgb) == "FFFFFF"
    assert common.text == "Witch's Hat" and not common.font.italic
    assert sci._r.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}outerShdw") is not None
    # lower-left with a small margin, no filled rectangle behind it
    assert box.left < Emu(914400 * 0.5)
    assert box.top + box.height > prs.slide_height - Emu(914400 * 0.5)
    assert box.top + box.height < prs.slide_height
    from pptx.enum.dml import MSO_FILL
    assert box.fill.type in (None, MSO_FILL.BACKGROUND)  # transparent: no box behind the text
    # Life-only observation: no annotation box at all
    assert not any(s.has_text_frame and s.text_frame.text.strip() for s in prs.slides[4].shapes)


def test_favorites_never_in_pptx(tmp_path):
    pr, obs = _deck({"annotations": AnnotationSettings(common_name=True, date=True, location=True, observer=True)})
    _, out, _ = _build(tmp_path, pr, obs)
    with zipfile.ZipFile(out) as z:
        xml = "".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    assert "♥" not in xml and "fave" not in xml.lower() and "favorite" not in xml.lower()


def test_divider_slides(tmp_path):
    pr, obs = _deck({"grouping": "genus", "dividers": True})
    prs, _, plan = _build(tmp_path, pr, obs)
    dividers = [(i, s) for i, s in enumerate(plan["slides"]) if s["kind"] == "divider"]
    assert [d[1]["text"] for d in dividers] == ["Hygrocybe", "Unclassified", "Amanita"]
    slide = prs.slides[dividers[0][0]]
    assert not _pictures(slide)
    texts = [p.runs[0] for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.runs]
    assert [t.text for t in texts] == ["GENUS", "Hygrocybe"]
    assert texts[1].font.italic is True and texts[1].font.size.pt >= 44


def test_title_slide_with_cropped_background(tmp_path):
    pr, obs = _deck({"title": "Mushrooms of Point Reyes", "presenter": "Alan Rockefeller",
                     "title_photo": TitlePhoto(observation_id=1, photo_id=102)})
    src = _jpeg(tmp_path / "title-src.jpg", 2000, 3000, color=(200, 180, 60))
    title = title_background(src, tmp_path / "title.jpg")
    with Image.open(title.path) as im:
        assert abs(im.width / im.height - 16 / 9) < 0.01  # cover-cropped to 16:9
        assert im.width <= 2000  # never upscaled past the source
        r, g, b = im.getpixel((im.width // 2, im.height // 2))
        assert max(r, g, b) < 120 and (max(r, g, b) - min(r, g, b)) < 80  # darkened, desaturated
    plan = build_slide_plan(pr, obs)
    assert plan["slides"][0]["photo"]["photo_id"] == 102
    out = tmp_path / "deck.pptx"
    write_pptx(plan, _images(tmp_path, obs), title, out)
    prs = Presentation(str(out))
    slide = prs.slides[0]
    (pic,) = _pictures(slide)
    assert (pic.left, pic.top, pic.width, pic.height) == (0, 0, prs.slide_width, prs.slide_height)
    texts = [p.runs[0] for s in slide.shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.runs]
    assert [t.text for t in texts] == ["Mushrooms of Point Reyes", "Alan Rockefeller"]
    assert texts[0].font.size.pt >= 44 and texts[0].font.bold
    assert texts[1].font.size.pt < texts[0].font.size.pt
    title_box = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "Mushrooms of Point Reyes")
    pres_box = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "Alan Rockefeller")
    assert pres_box.top > title_box.top  # presenter beneath the title


def test_title_slide_without_photo_on_white(tmp_path):
    pr, obs = _deck({"title": "Club Night", "background": "white"})
    prs, _, _ = _build(tmp_path, pr, obs)
    slide = prs.slides[0]
    assert not _pictures(slide)
    run = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "Club Night").text_frame.paragraphs[0].runs[0]
    assert str(run.font.color.rgb) != "FFFFFF"


def test_exif_rotation_is_applied(tmp_path):
    src = _jpeg(tmp_path / "rot.jpg", 400, 300, exif_orientation=6)  # rotate 90° on display
    prepared = prepare_for_slide(src, tmp_path / "w")
    assert (prepared.width, prepared.height) == (300, 400)
    assert prepared.path != src
    plain = prepare_for_slide(_jpeg(tmp_path / "plain.jpg", 400, 300), tmp_path / "w")
    assert plain.path.name == "plain.jpg"  # untouched


def test_speaker_notes_have_link_not_faves(tmp_path):
    pr, obs = _deck()
    prs, _, _ = _build(tmp_path, pr, obs)
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "https://www.inaturalist.org/observations/1" in notes and "CC BY" in notes
    assert "17" not in notes
