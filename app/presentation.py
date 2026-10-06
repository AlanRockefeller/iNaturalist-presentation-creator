"""Slide plan and PowerPoint generation.

``build_slide_plan`` decides what the deck contains. The browser's Preview and
``write_pptx`` both render that plan, so the preview shows the real slide
sequence.

Generation rules:
* 16:9 widescreen (13.333 x 7.5 in).
* Image slides never crop: the photo is scaled to fit and centred; the rest of
  the slide shows the chosen background (black or white).
* Photos stay picture objects; text stays editable text boxes.
* The title slide is the only slide where a photo is cropped (full bleed,
  darkened and desaturated for legibility).
* Favorite counts are never written anywhere in the deck.
* Large decks: image bytes are streamed from disk when the .pptx is written
  instead of being held in memory, and media are stored uncompressed (JPEGs do
  not deflate) so writing is fast.
"""

from __future__ import annotations

import hashlib
import zipfile
from datetime import date
from pathlib import Path
from typing import Callable

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.opc import serialized as _serialized
from pptx.oxml.ns import qn
from pptx.parts.image import ImagePart
from pptx.util import Emu, Pt

from .models import ObservationState, Project
from .project import observer_enabled
from .sorting import GROUP_LABELS, effective_name, group_of, passes_filters

SLIDE_W = Emu(12192000)  # 13.333 in
SLIDE_H = Emu(6858000)   # 7.5 in
EMU_PER_INCH = 914400

SCIENTIFIC_PT = 28
DETAIL_PT = 18
MARGIN_X = Emu(int(0.30 * EMU_PER_INCH))
MARGIN_Y = Emu(int(0.22 * EMU_PER_INCH))
TITLE_FONT = "Georgia"
BODY_FONT = "Calibri"
GOLD = RGBColor(0xC9, 0xA9, 0x62)

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]


def format_date(iso: str | None) -> str:
    """'2026-11-14' -> 'November 14, 2026'."""
    if not iso:
        return ""
    try:
        d = date.fromisoformat(iso[:10])
    except ValueError:
        return ""
    return f"{MONTHS[d.month - 1]} {d.day}, {d.year}"


def observer_name(obs: dict) -> str:
    user = obs.get("user") or {}
    return (user.get("name") or user.get("login") or "").strip()


def annotation_lines(obs: dict, state: ObservationState, project: Project) -> list[dict]:
    """Text lines for an image slide, in display order.

    Each line: {"field", "text", "italic", "size"}. Only the scientific name is
    italic. ``state.show_lines`` turns a line on or off for this observation
    regardless of the settings; an override of "" also hides it.
    """
    ann = project.settings.annotations
    ov = state.overrides
    lines = []

    def add(field: str, enabled: bool, override, default: str, italic: bool, size: int):
        enabled = state.show_lines.get(field, enabled)  # per-observation choice wins
        if not enabled:
            return
        text = default if override is None else override
        text = (text or "").strip()
        if text:
            lines.append({"field": field, "text": text, "italic": italic, "size": size})

    add("scientific", ann.scientific, ov.scientific, obs.get("inat_name") or "", True, SCIENTIFIC_PT)
    add("common_name", ann.common_name, ov.common_name, obs.get("common_name") or "", False, DETAIL_PT)
    add("date", ann.date, ov.date, format_date(obs.get("observed_on")), False, DETAIL_PT)
    add("location", ann.location, ov.location, obs.get("place_guess") or "", False, DETAIL_PT)
    name = observer_name(obs)
    add("observer", observer_enabled(project), ov.observer, f"Photo: {name}" if name else "", False, DETAIL_PT)
    return lines


def photo_positions(obs: dict, state: ObservationState) -> dict[int, int]:
    """Photo id -> slide position: ``state.photo_order`` first, then the rest in
    iNaturalist's order. Mirrored by orderedPhotos() in app.js."""
    ids = [p["id"] for p in obs.get("photos", [])]
    present = set(ids)
    ordered = [pid for pid in state.photo_order if pid in present]
    chosen = set(ordered)
    ordered += [pid for pid in ids if pid not in chosen]
    return {pid: n for n, pid in enumerate(ordered)}


def build_slide_plan(project: Project, workspace_obs: dict[int, dict], max_image_slides: int | None = None) -> dict:
    """The exact slide sequence, plus warnings and counts."""
    settings = project.settings
    states = {s.id: s for s in project.observations}
    slides: list[dict] = []
    warnings: list[str] = []

    title_photo = None
    tp = settings.title_photo
    if tp:
        obs = workspace_obs.get(tp.observation_id)
        photo = next((p for p in (obs or {}).get("photos", []) if p["id"] == tp.photo_id), None)
        if photo:
            tp_state = states.get(tp.observation_id)
            title_photo = {"observation_id": tp.observation_id, "photo_id": tp.photo_id,
                           "url": photo["url"], "position": tp.position,
                           "rotation": tp_state.rotations.get(tp.photo_id, 0) if tp_state else 0}
        else:
            warnings.append("The title background photo is no longer available on iNaturalist.")
    slides.append({
        "kind": "title",
        "title": settings.title.strip() or "Untitled presentation",
        "presenter": settings.presenter.strip(),
        "photo": title_photo,
    })

    grouping = settings.grouping
    current_group = None
    image_count = 0
    divider_count = 0
    unavailable_obs = 0
    unavailable_photos = 0
    observations_used = 0
    for oid in project.order:
        st = states.get(oid)
        if st is None:
            continue
        obs = workspace_obs.get(oid)
        if st.status == "unavailable" or obs is None:
            if st.selected_photo_ids:
                unavailable_obs += 1
            continue
        if not passes_filters(obs, st, project):
            continue
        by_id = {p["id"]: p for p in obs.get("photos", [])}
        photos = []
        for pid in st.selected_photo_ids:
            if pid in by_id:
                photos.append(by_id[pid])
            else:
                unavailable_photos += 1
        if not photos:
            continue
        # The user's photo order if they set one, then iNaturalist's order.
        position = photo_positions(obs, st)
        photos.sort(key=lambda p: position[p["id"]])

        if grouping != "none" and settings.dividers:
            _, label = group_of(obs, st, grouping, project)
            if label != current_group:
                current_group = label
                slides.append({
                    "kind": "divider",
                    "text": label,
                    "rank_label": GROUP_LABELS.get(grouping, ""),
                    "italic": grouping == "genus" and label != "Unclassified",
                })
                divider_count += 1
        lines = annotation_lines(obs, st, project)
        observations_used += 1
        for p in photos:
            rotation = st.rotations.get(p["id"], 0)
            width, height = p.get("width"), p.get("height")
            if rotation in (90, 270):
                width, height = height, width
            slides.append({
                "kind": "image",
                "observation_id": oid,
                "photo_id": p["id"],
                "url": p["url"],
                "rotation": rotation,
                "width": width,
                "height": height,
                "lines": lines,
                "notes": _speaker_notes(obs, st, p) if settings.speaker_notes else "",
            })
            image_count += 1

    if unavailable_obs:
        warnings.append(f"{unavailable_obs} selected observation(s) are no longer available on iNaturalist and are skipped.")
    if unavailable_photos:
        warnings.append(f"{unavailable_photos} selected photo(s) were deleted from iNaturalist and are skipped.")
    if max_image_slides is not None and image_count > max_image_slides:
        warnings.append(f"This presentation has {image_count:,} photo slides; the limit is {max_image_slides:,}.")
    return {
        "slides": slides,
        "counts": {
            "total": len(slides), "title": 1, "dividers": divider_count, "images": image_count,
            "observations": observations_used,
        },
        "warnings": warnings,
        "background": settings.background,
    }


def _speaker_notes(obs: dict, st: ObservationState, photo: dict) -> str:
    bits = [effective_name(st, obs) or "Unidentified"]
    if obs.get("common_name"):
        bits.append(obs["common_name"])
    bits.append(obs.get("uri") or "")
    if photo.get("attribution"):
        bits.append(photo["attribution"])
    return "\n".join(b for b in bits if b)


# ---------------------------------------------------------------------------
# PPTX writing
# ---------------------------------------------------------------------------

class FileImagePart(ImagePart):
    """An image part whose bytes stay on disk until the package is written."""

    def __init__(self, partname, content_type, package, path: Path, on_write=None):
        super().__init__(partname, content_type, package, b"", path.name)
        self._path = path
        self._on_write = on_write

    @property
    def blob(self) -> bytes:  # read lazily, one image at a time
        data = self._path.read_bytes()
        if self._on_write:
            self._on_write()
        return data

    @property
    def sha1(self) -> str:  # unique per file; avoids hashing every image on add
        return hashlib.sha1(str(self._path).encode()).hexdigest()

    def scale(self, scaled_cx, scaled_cy):
        # Callers always pass an explicit size; the base class would decode the
        # image just to compute a native size it then discards.
        if not (scaled_cx and scaled_cy):
            raise ValueError("FileImagePart needs an explicit width and height")
        return scaled_cx, scaled_cy


_orig_zip_write = _serialized._ZipPkgWriter.write


def _zip_write(self, pack_uri, blob):
    # Media is already compressed; deflating JPEGs wastes minutes on big decks.
    if pack_uri.membername.startswith("ppt/media/"):
        info = zipfile.ZipInfo(pack_uri.membername, date_time=(2026, 1, 1, 0, 0, 0))
        info.compress_type = zipfile.ZIP_STORED
        self._zipf.writestr(info, blob)
    else:
        _orig_zip_write(self, pack_uri, blob)


_serialized._ZipPkgWriter.write = _zip_write


def _rgb(background: str) -> RGBColor:
    return RGBColor(0, 0, 0) if background == "black" else RGBColor(0xFF, 0xFF, 0xFF)


def _text_colors(background: str) -> tuple[RGBColor, str]:
    """Annotation color and shadow color. White text with a dark shadow on
    black; on white slides the letterbox is white, so dark text with a light halo."""
    if background == "black":
        return RGBColor(0xFF, 0xFF, 0xFF), "000000"
    return RGBColor(0x11, 0x11, 0x11), "FFFFFF"


def _add_shadow(run, color_hex: str) -> None:
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(qn("a:effectLst")):
        rPr.remove(old)
    effect = etree.SubElement(rPr, qn("a:effectLst"))
    shadow = etree.SubElement(effect, qn("a:outerShdw"), {
        "blurRad": "50800", "dist": "19050", "dir": "2700000", "algn": "tl", "rotWithShape": "0",
    })
    clr = etree.SubElement(shadow, qn("a:srgbClr"), {"val": color_hex})
    etree.SubElement(clr, qn("a:alpha"), {"val": "75000"})
    # effectLst must precede latin/ea/cs font elements in CT_TextCharacterProperties
    for tag in ("a:latin", "a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
            rPr.append(el)


def _set_background(slide, background: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(background)


def _textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.BOTTOM):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return box, tf


def _add_para(tf, first: bool, text: str, size: int, italic=False, bold=False, color=None,
              font=BODY_FONT, align=PP_ALIGN.LEFT, shadow: str | None = None, spacing: int | None = None):
    para = tf.paragraphs[0] if first else tf.add_paragraph()
    para.alignment = align
    run = para.add_run()
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.italic = italic
    f.bold = bold
    f.name = font
    if color is not None:
        f.color.rgb = color
    if spacing is not None:
        run._r.get_or_add_rPr().set("spc", str(spacing))
    if shadow:
        _add_shadow(run, shadow)
    return para


def fit_rect(img_w: int, img_h: int, box_w: int = int(SLIDE_W), box_h: int = int(SLIDE_H)) -> tuple[int, int, int, int]:
    """Largest (left, top, width, height) with the image's aspect ratio inside the
    box, centred. Never crops."""
    if img_w <= 0 or img_h <= 0:
        raise ValueError("image has no size")
    scale = min(box_w / img_w, box_h / img_h)
    w = int(round(img_w * scale))
    h = int(round(img_h * scale))
    w, h = min(w, box_w), min(h, box_h)
    return (box_w - w) // 2, (box_h - h) // 2, w, h


class DeckWriter:
    def __init__(self, background: str, on_media_write: Callable[[int], None] | None = None):
        self._on_media_write = on_media_write
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.layout = self.prs.slide_layouts[6]  # blank
        self.background = background
        self._media_n = 0
        self.media_written = 0

    def _new_slide(self):
        slide = self.prs.slides.add_slide(self.layout)
        _set_background(slide, self.background)
        return slide

    def _picture(self, slide, path: Path, content_type: str, ext: str, left, top, width, height):
        self._media_n += 1
        package = self.prs.part.package
        partname = PackURI(f"/ppt/media/photo{self._media_n}.{ext}")
        part = FileImagePart(partname, content_type, package, path, on_write=self._count_write)
        rId = slide.part.relate_to(part, RT.IMAGE)
        return slide.shapes._add_pic_from_image_part(part, rId, left, top, width, height)

    def _count_write(self):
        self.media_written += 1
        if self._on_media_write:
            self._on_media_write(self.media_written)

    def add_title(self, title: str, presenter: str, bg_image=None):
        slide = self._new_slide()
        light_text = self.background == "black" or bg_image is not None
        if bg_image is not None:
            self._picture(slide, bg_image.path, bg_image.content_type, bg_image.ext, 0, 0, SLIDE_W, SLIDE_H)
        color = RGBColor(0xFF, 0xFF, 0xFF) if light_text else RGBColor(0x1A, 0x2F, 0x23)
        shadow = "000000" if bg_image is not None else None
        w = int(SLIDE_W * 0.84)
        left = (int(SLIDE_W) - w) // 2
        box, tf = _textbox(slide, left, int(SLIDE_H * 0.16), w, int(SLIDE_H * 0.42), MSO_ANCHOR.BOTTOM)
        _add_para(tf, True, title, 54, bold=True, color=color, font=TITLE_FONT, align=PP_ALIGN.CENTER, shadow=shadow)
        rule_w = int(SLIDE_W * 0.12)
        rule = slide.shapes.add_shape(1, (int(SLIDE_W) - rule_w) // 2, int(SLIDE_H * 0.625), rule_w, Emu(28575))
        rule.fill.solid()
        rule.fill.fore_color.rgb = GOLD
        rule.line.fill.background()
        if presenter:
            box, tf = _textbox(slide, left, int(SLIDE_H * 0.67), w, int(SLIDE_H * 0.2), MSO_ANCHOR.TOP)
            _add_para(tf, True, presenter, 28, color=color, font=TITLE_FONT, align=PP_ALIGN.CENTER, shadow=shadow)
        return slide

    def add_divider(self, text: str, rank_label: str, italic: bool):
        slide = self._new_slide()
        color = RGBColor(0xFF, 0xFF, 0xFF) if self.background == "black" else RGBColor(0x1A, 0x2F, 0x23)
        w = int(SLIDE_W * 0.84)
        left = (int(SLIDE_W) - w) // 2
        if rank_label:
            box, tf = _textbox(slide, left, int(SLIDE_H * 0.30), w, int(SLIDE_H * 0.1), MSO_ANCHOR.BOTTOM)
            _add_para(tf, True, rank_label.upper(), 16, color=GOLD, font=BODY_FONT, align=PP_ALIGN.CENTER, spacing=300)
        box, tf = _textbox(slide, left, int(SLIDE_H * 0.42), w, int(SLIDE_H * 0.22), MSO_ANCHOR.TOP)
        _add_para(tf, True, text, 54, italic=italic, color=color, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        return slide

    def add_image(self, prepared, lines: list[dict], notes: str = "", rotation: int = 0):
        slide = self._new_slide()
        rotation %= 360
        if rotation in (90, 270):
            # Fit the turned photo, then size the frame as the unturned photo:
            # PowerPoint rotates a picture about its centre.
            left, top, w, h = fit_rect(prepared.height, prepared.width)
            cx, cy = left + w // 2, top + h // 2
            left, top, w, h = cx - h // 2, cy - w // 2, h, w
        else:
            left, top, w, h = fit_rect(prepared.width, prepared.height)
        pic = self._picture(slide, prepared.path, prepared.content_type, prepared.ext, left, top, w, h)
        if rotation:
            # Rotating the picture keeps the original bytes; nothing is re-encoded.
            pic.rot = float(rotation)  # the <p:pic> element; same as Picture.rotation
        if lines:
            color, shadow = _text_colors(self.background)
            box_w = int(SLIDE_W * 0.8)
            box_h = int(SLIDE_H * 0.4)
            box, tf = _textbox(slide, MARGIN_X, int(SLIDE_H) - int(MARGIN_Y) - box_h, box_w, box_h, MSO_ANCHOR.BOTTOM)
            box.name = "Annotation"
            for i, line in enumerate(lines):
                _add_para(tf, i == 0, line["text"], line["size"], italic=line["italic"], color=color, shadow=shadow)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def save(self, out_path: Path, title: str = "", author: str = "") -> None:
        props = self.prs.core_properties
        props.title = title[:200]
        props.author = author[:200] or "Dikarya Presentations"
        props.comments = "Created with presentations.dikarya.us from iNaturalist observations."
        self.prs.save(str(out_path))


def write_pptx(
    plan: dict,
    images: dict[int, object],
    title_image,
    out_path: Path,
    progress: Callable[[int, int], None] | None = None,
    author: str = "",
    save_progress: Callable[[int, int], None] | None = None,
) -> dict:
    """Render a slide plan. ``images`` maps photo id -> PreparedImage; slides
    whose photo is missing are skipped and counted."""
    media_total = [0]
    writer = DeckWriter(
        plan.get("background", "black"),
        (lambda n: save_progress(n, media_total[0])) if save_progress else None,
    )
    skipped = 0
    slides = plan["slides"]
    total = len(slides)
    title_text = ""
    for n, spec in enumerate(slides, 1):
        kind = spec["kind"]
        if kind == "title":
            title_text = spec["title"]
            writer.add_title(spec["title"], spec.get("presenter", ""), title_image)
        elif kind == "divider":
            writer.add_divider(spec["text"], spec.get("rank_label", ""), spec.get("italic", False))
        elif kind == "image":
            prepared = images.get(spec["photo_id"])
            if prepared is None:
                skipped += 1
            else:
                writer.add_image(prepared, spec.get("lines", []), spec.get("notes", ""), spec.get("rotation", 0))
        if progress:
            progress(n, total)
    media_total[0] = writer._media_n
    writer.save(out_path, title=title_text, author=author)
    return {"slides": len(writer.prs.slides), "skipped": skipped}
